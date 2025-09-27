import json
import torch
import whisper
import torchaudio
import argparse
import uuid
import os
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Any, TypedDict
from docx import Document
from dotenv import load_dotenv

from pyannote.audio import Pipeline
from langchain_anthropic import ChatAnthropic
from langchain_core.messages import HumanMessage, SystemMessage
from langgraph.graph import StateGraph, END
from langgraph.checkpoint.memory import MemorySaver

# Load environment variables from .env file
load_dotenv()

# State definition for LangGraph
class TranscriptionState(TypedDict):
    audio_file_path: str
    auth_token: str
    whisper_model_size: str
    session_id: str
    output_dir: str
    original_transcript: List[Dict]
    named_transcript: List[Dict]
    verification_result: bool
    issues_found: List[Dict]
    final_transcript: List[Dict]
    retry_count: int
    max_retries: int

class AudioTranscriptionPipeline:
    def __init__(self, claude_api_key: str, max_retries: int = 3):
        """
        Initialize the transcription pipeline with LangGraph
        
        Args:
            claude_api_key: Anthropic API key for Claude
            max_retries: Maximum number of retry attempts for fixing issues
        """
        self.claude_llm = ChatAnthropic(
            model="claude-sonnet-4-20250514",
            api_key=claude_api_key,
            temperature=0.1
        )
        self.max_retries = max_retries
        
        # Build the LangGraph workflow
        self.workflow = self._build_workflow()
        self.memory = MemorySaver()
        self.app = self.workflow.compile(checkpointer=self.memory)
    
    def _create_session_directory(self, audio_file_path: str) -> tuple[str, str]:
        """
        Create a unique session directory for this transcription run
        
        Args:
            audio_file_path: Path to the audio file
            
        Returns:
            tuple: (session_id, output_directory_path)
        """
        # Generate unique session ID
        session_uuid = str(uuid.uuid4())[:8]  # Short UUID
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        audio_filename = Path(audio_file_path).stem  # Get filename without extension
        
        session_id = f"{timestamp}_{audio_filename}_{session_uuid}"
        
        # Create directory structure
        base_dir = Path("data")
        output_dir = base_dir / session_id
        output_dir.mkdir(parents=True, exist_ok=True)
        
        # Create subdirectories
        (output_dir / "transcripts").mkdir(exist_ok=True)
        (output_dir / "logs").mkdir(exist_ok=True)
        (output_dir / "outputs").mkdir(exist_ok=True)
        
        print(f"📁 Created session directory: {output_dir}")
        
        # Create session info file
        session_info = {
            "session_id": session_id,
            "created_at": datetime.now().isoformat(),
            "audio_file": audio_file_path,
            "audio_filename": audio_filename,
            "status": "started"
        }
        
        with open(output_dir / "session_info.json", 'w') as f:
            json.dump(session_info, f, indent=2)
        
        return session_id, str(output_dir)
    
    def _log_step(self, output_dir: str, step_name: str, data: Any, status: str = "completed"):
        """Log step information"""
        log_entry = {
            "step": step_name,
            "timestamp": datetime.now().isoformat(),
            "status": status,
            "data_size": len(data) if isinstance(data, (list, dict)) else str(type(data))
        }
        
        log_file = Path(output_dir) / "logs" / "pipeline.log"
        
        # Append to log file
        with open(log_file, 'a') as f:
            f.write(json.dumps(log_entry) + "\n")
    
    def _build_workflow(self) -> StateGraph:
        """Build the LangGraph workflow"""
        workflow = StateGraph(TranscriptionState)
        
        # Add nodes
        workflow.add_node("setup_session", self._setup_session)
        workflow.add_node("transcribe_audio", self._transcribe_audio)
        workflow.add_node("name_speakers", self._name_speakers)
        workflow.add_node("verify_transcript", self._verify_transcript)
        workflow.add_node("fix_issues", self._fix_issues)
        workflow.add_node("generate_outputs", self._generate_outputs)
        workflow.add_node("finalize_session", self._finalize_session)
        
        # Define the flow
        workflow.set_entry_point("setup_session")
        workflow.add_edge("setup_session", "transcribe_audio")
        workflow.add_edge("transcribe_audio", "name_speakers")
        workflow.add_edge("name_speakers", "verify_transcript")
        
        # Conditional routing based on verification
        workflow.add_conditional_edges(
            "verify_transcript",
            self._should_fix_or_continue,
            {
                "fix": "fix_issues",
                "continue": "generate_outputs"
            }
        )
        
        # After fixing, go back to verification unless max retries reached
        workflow.add_conditional_edges(
            "fix_issues",
            self._should_retry_or_continue,
            {
                "retry": "verify_transcript",
                "continue": "generate_outputs"
            }
        )
        
        workflow.add_edge("generate_outputs", "finalize_session")
        workflow.add_edge("finalize_session", END)
        
        return workflow
    
    def _setup_session(self, state: TranscriptionState) -> TranscriptionState:
        """Step 0: Setup session directory and initialize"""
        print("🚀 Setting up transcription session...")
        
        session_id, output_dir = self._create_session_directory(state["audio_file_path"])
        
        state["session_id"] = session_id
        state["output_dir"] = output_dir
        state["retry_count"] = 0
        state["max_retries"] = self.max_retries
        
        self._log_step(output_dir, "setup_session", {"session_id": session_id}, "started")
        
        return state
    
    def _transcribe_audio(self, state: TranscriptionState) -> TranscriptionState:
        """Step 1: Transcribe audio with speaker diarization"""
        print("🎵 Starting audio transcription and diarization...")
        
        output_dir = Path(state["output_dir"])
        
        # Initialize pipeline
        pipeline = Pipeline.from_pretrained(
            "pyannote/speaker-diarization-3.1",
            use_auth_token=state["auth_token"]
        )
        
        # Device setup
        device = torch.device("mps") if torch.backends.mps.is_available() else torch.device("cpu")
        print(f"Using device: {device}")
        pipeline.to(device)
        
        # Load models and audio
        asr_model = whisper.load_model(state["whisper_model_size"], device="cpu")
        waveform, sr = torchaudio.load(state["audio_file_path"])
        diarization = pipeline(state["audio_file_path"])
        
        # Process audio segments
        speaker_map = {}
        speaker_counter = 1
        transcript_segments = []
        
        for turn, _, speaker in diarization.itertracks(yield_label=True):
            # Map speakers
            if speaker not in speaker_map:
                speaker_map[speaker] = f"Speaker_{speaker_counter:02d}"
                speaker_counter += 1
            
            speaker_label = speaker_map[speaker]
            
            # Process audio segment
            start_frame = int(turn.start * sr)
            end_frame = int(turn.end * sr)
            segment_waveform = waveform[:, start_frame:end_frame]
            
            # Save and transcribe segment
            temp_file = output_dir / "temp_segment.wav"
            torchaudio.save(str(temp_file), segment_waveform, sr)
            result = asr_model.transcribe(str(temp_file), fp16=False)
            text = result["text"].strip()
            
            # Clean up temp file
            if temp_file.exists():
                temp_file.unlink()
            
            # Store segment info
            segment = {
                'start': turn.start,
                'end': turn.end,
                'speaker': speaker_label,
                'text': text,
                'segment_id': len(transcript_segments)
            }
            transcript_segments.append(segment)
            
            print(f"[{turn.start:.1f}s - {turn.end:.1f}s] {speaker_label}: {text}")
        
        # Save original transcript to JSON
        original_file = output_dir / "transcripts" / "01_original_transcript.json"
        
        with open(original_file, 'w', encoding='utf-8') as f:
            json.dump(transcript_segments, f, indent=2, ensure_ascii=False)
        
        print(f"✅ Original transcript saved to: {original_file}")
        
        # Update state
        state["original_transcript"] = transcript_segments
        
        # Log step
        self._log_step(state["output_dir"], "transcribe_audio", transcript_segments)
        
        return state
    
    def _name_speakers(self, state: TranscriptionState) -> TranscriptionState:
        """Step 2: Use LLM to identify and name speakers from context"""
        print("🤖 Identifying speaker names from context...")
        
        output_dir = Path(state["output_dir"])
        transcript_segments = state["original_transcript"]
        
        # Prepare context for LLM
        context_text = "\n".join([
            f"[{seg['start']:.1f}s-{seg['end']:.1f}s] {seg['speaker']}: {seg['text']}"
            for seg in transcript_segments
        ])
        
        system_prompt = """You are an expert at identifying speakers in transcripts by analyzing context clues, introductions, and conversational patterns.

Your task is to:
1. Analyze the transcript for names mentioned in the conversation
2. Identify which speaker corresponds to which name based on context
3. Replace generic speaker labels (like Speaker_01, Speaker_02) with actual names when possible
4. Maintain speaker consistency throughout the transcript
5. If you can't identify a name with high confidence, keep the original speaker label

Return a JSON object with speaker mappings in this format:
{
  "speaker_mappings": {
    "Speaker_01": "John Smith",
    "Speaker_02": "Speaker_02"
  },
  "confidence_notes": "Brief explanation of how names were identified",
  "names_found": ["John Smith"],
  "unnamed_speakers": ["Speaker_02"]
}"""

        messages = [
            SystemMessage(content=system_prompt),
            HumanMessage(content=f"Please analyze this transcript and identify speaker names:\n\n{context_text}")
        ]
        
        response = self.claude_llm.invoke(messages)
        
        try:
            # Parse LLM response
            response_text = response.content
            if "```json" in response_text:
                json_part = response_text.split("```json")[1].split("```")[0]
            else:
                json_part = response_text
            
            speaker_info = json.loads(json_part)
            speaker_mappings = speaker_info.get("speaker_mappings", {})
            
            # Save speaker analysis
            analysis_file = output_dir / "logs" / "speaker_analysis.json"
            with open(analysis_file, 'w') as f:
                json.dump(speaker_info, f, indent=2)
            
            # Apply speaker name mappings
            named_segments = []
            for segment in transcript_segments:
                new_segment = segment.copy()
                original_speaker = segment["speaker"]
                new_segment["speaker"] = speaker_mappings.get(original_speaker, original_speaker)
                new_segment["original_speaker"] = original_speaker
                named_segments.append(new_segment)
            
            # Save named transcript
            named_file = output_dir / "transcripts" / "02_named_transcript.json"
            
            with open(named_file, 'w', encoding='utf-8') as f:
                json.dump(named_segments, f, indent=2, ensure_ascii=False)
            
            print(f"✅ Named transcript saved to: {named_file}")
            print(f"Speaker mappings: {speaker_mappings}")
            
            state["named_transcript"] = named_segments
            
        except Exception as e:
            print(f"❌ Error in speaker naming: {e}")
            # Fallback to original transcript
            state["named_transcript"] = transcript_segments
        
        # Log step
        self._log_step(state["output_dir"], "name_speakers", state["named_transcript"])
        
        return state
    
    def _verify_transcript(self, state: TranscriptionState) -> TranscriptionState:
        """Step 3: Verify that speaker naming didn't alter content"""
        print("🔍 Verifying transcript integrity...")
        
        output_dir = Path(state["output_dir"])
        original_segments = state["original_transcript"]
        named_segments = state["named_transcript"]
        
        # Prepare comparison for LLM
        comparison_data = []
        for i, (orig, named) in enumerate(zip(original_segments, named_segments)):
            comparison_data.append({
                "segment_id": i,
                "original": {
                    "speaker": orig["speaker"],
                    "text": orig["text"],
                    "timing": f"{orig['start']:.1f}s-{orig['end']:.1f}s"
                },
                "named": {
                    "speaker": named["speaker"],
                    "text": named["text"],
                    "timing": f"{named['start']:.1f}s-{named['end']:.1f}s"
                }
            })
        
        system_prompt = """You are a transcript verification expert. Your job is to ensure that when speaker labels were updated, no content was accidentally modified.

Compare the original and named transcripts and check for:
1. Any changes to the actual spoken text content
2. Any changes to timestamps
3. Any missing or added segments
4. Speaker label consistency issues

The ONLY acceptable changes are speaker label updates (e.g., "Speaker_01" → "John Smith").

Return a JSON object with this format:
{
  "verification_passed": true/false,
  "issues_found": [
    {
      "segment_id": 0,
      "issue_type": "text_change",
      "description": "Text content was modified",
      "original_text": "...",
      "named_text": "...",
      "severity": "high"
    }
  ],
  "summary": "Brief summary of verification results",
  "stats": {
    "total_segments": 0,
    "segments_with_issues": 0,
    "acceptable_changes": 0
  }
}"""

        messages = [
            SystemMessage(content=system_prompt),
            HumanMessage(content=f"Please verify these transcripts:\n\n{json.dumps(comparison_data, indent=2)}")
        ]
        
        response = self.claude_llm.invoke(messages)
        
        try:
            response_text = response.content
            if "```json" in response_text:
                json_part = response_text.split("```json")[1].split("```")[0]
            else:
                json_part = response_text
            
            verification_result = json.loads(json_part)
            
            passed = verification_result.get("verification_passed", False)
            issues = verification_result.get("issues_found", [])
            
            # Save verification results
            verification_file = output_dir / "logs" / f"verification_attempt_{state['retry_count'] + 1}.json"
            with open(verification_file, 'w') as f:
                json.dump(verification_result, f, indent=2)
            
            print(f"Verification result: {'✅ PASSED' if passed else '❌ FAILED'}")
            if issues:
                print(f"Issues found: {len(issues)}")
                # Save issues to main issues file
                issues_file = output_dir / "logs" / "issues_found.json"
                with open(issues_file, 'w') as f:
                    json.dump(issues, f, indent=2)
            
            state["verification_result"] = passed
            state["issues_found"] = issues
            
        except Exception as e:
            print(f"❌ Error in verification: {e}")
            state["verification_result"] = False
            state["issues_found"] = [{"error": str(e)}]
        
        # Log step
        self._log_step(state["output_dir"], "verify_transcript", 
                      {"passed": state["verification_result"], "issues_count": len(state["issues_found"])})
        
        return state
    
    def _fix_issues(self, state: TranscriptionState) -> TranscriptionState:
        """Step 4: Fix identified issues"""
        print("🔧 Fixing identified issues...")
        
        output_dir = Path(state["output_dir"])
        issues = state["issues_found"]
        named_segments = state["named_transcript"]
        original_segments = state["original_transcript"]
        
        system_prompt = """You are a transcript repair expert. Given a list of issues found in a transcript, fix them by restoring the original content while preserving the correct speaker names.

Rules for fixing:
1. Restore original text content exactly as it was
2. Restore original timestamps exactly as they were
3. Keep the updated speaker names (the only acceptable changes)
4. Do not make any other modifications

Return the corrected transcript segments as a JSON array."""

        fix_context = {
            "issues": issues,
            "original_segments": original_segments,
            "named_segments": named_segments
        }
        
        messages = [
            SystemMessage(content=system_prompt),
            HumanMessage(content=f"Please fix these issues:\n\n{json.dumps(fix_context, indent=2)}")
        ]
        
        response = self.claude_llm.invoke(messages)
        
        try:
            response_text = response.content
            if "```json" in response_text:
                json_part = response_text.split("```json")[1].split("```")[0]
            else:
                json_part = response_text
            
            fixed_segments = json.loads(json_part)
            
            # Save fixed transcript
            fixed_file = output_dir / "transcripts" / f"03_fixed_transcript_attempt_{state['retry_count'] + 1}.json"
            with open(fixed_file, 'w') as f:
                json.dump(fixed_segments, f, indent=2)
            
            print(f"✅ Applied fixes to {len(issues)} issues")
            state["named_transcript"] = fixed_segments
            state["retry_count"] += 1
            
        except Exception as e:
            print(f"❌ Error in fixing issues: {e}")
            state["retry_count"] += 1
        
        # Log step
        self._log_step(state["output_dir"], "fix_issues", 
                      {"retry_count": state["retry_count"], "issues_fixed": len(issues)})
        
        return state
    
    def _generate_outputs(self, state: TranscriptionState) -> TranscriptionState:
        """Step 5: Generate all output files"""
        print("📄 Generating output files...")
        
        output_dir = Path(state["output_dir"])
        final_segments = state["named_transcript"]
        
        # Save final transcript JSON
        final_json = output_dir / "outputs" / "final_transcript.json"
        with open(final_json, 'w', encoding='utf-8') as f:
            json.dump(final_segments, f, indent=2, ensure_ascii=False)
        
        # Generate Word document
        doc = Document()
        
        # Add header
        doc.add_heading('Audio Transcript - Named Version', 0)
        doc.add_paragraph(f'Session ID: {state["session_id"]}')
        doc.add_paragraph(f'Source: {state["audio_file_path"]}')
        doc.add_paragraph(f'Generated on: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
        doc.add_paragraph(f'Verification: {"✅ PASSED" if state.get("verification_result", False) else "❌ FAILED"}')
        doc.add_paragraph(f'Total segments: {len(final_segments)}')
        
        if state["retry_count"] > 0:
            doc.add_paragraph(f'Retry attempts: {state["retry_count"]}')
        
        doc.add_paragraph('')  # Empty line
        
        # Add transcript content
        for segment in final_segments:
            time_stamp = f'[{segment["start"]:.1f}s - {segment["end"]:.1f}s]'
            speaker = segment["speaker"]
            text = segment["text"]
            
            paragraph = doc.add_paragraph()
            paragraph.add_run(f'{time_stamp} ').bold = True
            paragraph.add_run(f'{speaker}: ').italic = True
            paragraph.add_run(text)
        
        # Save Word document
        docx_file = output_dir / "outputs" / "final_transcript.docx"
        doc.save(str(docx_file))
        
        # Generate summary report
        summary = {
            "session_info": {
                "session_id": state["session_id"],
                "audio_file": state["audio_file_path"],
                "created_at": datetime.now().isoformat()
            },
            "processing_stats": {
                "total_segments": len(final_segments),
                "verification_passed": state["verification_result"],
                "retry_attempts": state["retry_count"],
                "issues_found": len(state["issues_found"])
            },
            "output_files": {
                "final_transcript_json": str(final_json),
                "final_transcript_docx": str(docx_file),
                "session_directory": str(output_dir)
            },
            "speakers_identified": list(set(seg["speaker"] for seg in final_segments))
        }
        
        summary_file = output_dir / "outputs" / "summary_report.json"
        with open(summary_file, 'w') as f:
            json.dump(summary, f, indent=2)
        
        print(f"✅ Final outputs generated:")
        print(f"   - Transcript (JSON): {final_json}")
        print(f"   - Transcript (DOCX): {docx_file}")
        print(f"   - Summary Report: {summary_file}")
        
        state["final_transcript"] = final_segments
        
        # Log step
        self._log_step(state["output_dir"], "generate_outputs", summary)
        
        return state
    
    def _finalize_session(self, state: TranscriptionState) -> TranscriptionState:
        """Step 6: Finalize session and update status"""
        print("🎯 Finalizing session...")
        
        output_dir = Path(state["output_dir"])
        
        # Update session info
        session_info_file = output_dir / "session_info.json"
        with open(session_info_file, 'r') as f:
            session_info = json.load(f)
        
        session_info.update({
            "completed_at": datetime.now().isoformat(),
            "status": "completed",
            "verification_passed": state["verification_result"],
            "retry_count": state["retry_count"],
            "total_segments": len(state["final_transcript"]),
            "speakers_found": list(set(seg["speaker"] for seg in state["final_transcript"]))
        })
        
        with open(session_info_file, 'w') as f:
            json.dump(session_info, f, indent=2)
        
        # Final log entry
        self._log_step(state["output_dir"], "finalize_session", session_info, "completed")
        
        print(f"✅ Session finalized: {state['session_id']}")
        
        return state
    
    def _should_fix_or_continue(self, state: TranscriptionState) -> str:
        """Conditional routing: fix issues or continue"""
        if state["verification_result"]:
            return "continue"
        else:
            return "fix"
    
    def _should_retry_or_continue(self, state: TranscriptionState) -> str:
        """Conditional routing: retry verification or continue"""
        if state["retry_count"] >= state["max_retries"]:
            print(f"⚠️ Maximum retries ({state['max_retries']}) reached. Proceeding with current result.")
            return "continue"
        else:
            return "retry"
    
    def process_audio_file(self, audio_file_path: str, auth_token: str, whisper_model_size: str = "small"):
        """
        Process audio file through the complete LangGraph pipeline
        
        Args:
            audio_file_path: Path to the audio file
            auth_token: HuggingFace authentication token
            whisper_model_size: Whisper model size to use
        """
        # Validate audio file exists
        if not Path(audio_file_path).exists():
            raise FileNotFoundError(f"Audio file not found: {audio_file_path}")
        
        # Initial state
        initial_state = {
            "audio_file_path": audio_file_path,
            "auth_token": auth_token,
            "whisper_model_size": whisper_model_size,
            "session_id": "",
            "output_dir": "",
            "original_transcript": [],
            "named_transcript": [],
            "verification_result": False,
            "issues_found": [],
            "final_transcript": [],
            "retry_count": 0,
            "max_retries": self.max_retries
        }
        
        print("🚀 Starting LangGraph Audio Transcription Pipeline")
        print("=" * 80)
        print(f"Audio File: {audio_file_path}")
        print(f"Whisper Model: {whisper_model_size}")
        print(f"Max Retries: {self.max_retries}")
        print("=" * 80)
        
        # Run the workflow
        config = {"configurable": {"thread_id": f"transcription_{datetime.now().strftime('%Y%m%d_%H%M%S')}"}}
        final_state = self.app.invoke(initial_state, config)
        
        print("=" * 80)
        print("✅ Pipeline completed successfully!")
        print(f"Session ID: {final_state['session_id']}")
        print(f"Output Directory: {final_state['output_dir']}")
        print(f"Verification Status: {'✅ PASSED' if final_state['verification_result'] else '❌ FAILED'}")
        print(f"Total Segments: {len(final_state['final_transcript'])}")
        print("=" * 80)
        
        return final_state

def load_environment_variables():
    """Load and validate environment variables"""
    env_vars = {
        'ANTHROPIC_API_KEY': os.getenv('ANTHROPIC_API_KEY'),
        'HUGGINGFACE_TOKEN': os.getenv('HUGGINGFACE_TOKEN')
    }
    
    # Optional LangGraph Studio visualization
    langchain_api_key = os.getenv('LANGCHAIN_API_KEY')
    if langchain_api_key:
        env_vars['LANGCHAIN_API_KEY'] = langchain_api_key
        print("🎨 LangGraph Studio visualization enabled")
    
    missing_vars = [var for var, value in env_vars.items() if not value and var in ['ANTHROPIC_API_KEY', 'HUGGINGFACE_TOKEN']]
    
    if missing_vars:
        print("❌ Missing required environment variables:")
        for var in missing_vars:
            print(f"   - {var}")
        print("\nPlease create a .env file with the following variables:")
        print("ANTHROPIC_API_KEY=your_claude_api_key_here")
        print("HUGGINGFACE_TOKEN=your_hf_token_here")
        print("\nOptional for LangGraph Studio visualization:")
        print("LANGCHAIN_API_KEY=your_langsmith_api_key_here")
        print("LANGCHAIN_TRACING_V2=true")
        print("LANGCHAIN_PROJECT=audio-transcription-pipeline")
        return None
    
    return env_vars

def get_audio_file_path(provided_path=None):
    """Get audio file path either from argument or interactive input"""
    if provided_path:
        return provided_path
    
    print("🎵 Audio Transcription Pipeline")
    print("=" * 40)
    
    while True:
        audio_file = input("📁 Please enter the path to your audio file: ").strip()
        
        if not audio_file:
            print("❌ Please provide a valid file path")
            continue
        
        # Handle quoted paths
        audio_file = audio_file.strip('"\'')
        
        if Path(audio_file).exists():
            return audio_file
        else:
            print(f"❌ File not found: {audio_file}")
            retry = input("🔄 Would you like to try again? (y/n): ").lower().strip()
            if retry not in ['y', 'yes']:
                print("👋 Exiting...")
                exit(1)

def main():
    """Command line interface"""
    parser = argparse.ArgumentParser(
        description="LangGraph Audio Transcription Pipeline",
        epilog="""
Environment Variables (required in .env file):
  ANTHROPIC_API_KEY     Your Anthropic Claude API key
  HUGGINGFACE_TOKEN     Your HuggingFace authentication token

Example .env file:
  ANTHROPIC_API_KEY=sk-ant-api03-...
  HUGGINGFACE_TOKEN=hf_...

Example usage:
  python transcription_pipeline.py                              # Interactive mode
  python transcription_pipeline.py audio.mp3                   # Direct file
  python transcription_pipeline.py /path/to/meeting.wav --whisper-model medium
        """,
        formatter_class=argparse.RawDescriptionHelpFormatter
    )
    
    parser.add_argument("audio_file", nargs='?', help="Path to the audio file (optional - will prompt if not provided)")
    parser.add_argument("--claude-api-key", 
                       help="Anthropic Claude API key (overrides .env file)")
    parser.add_argument("--hf-token", 
                       help="HuggingFace authentication token (overrides .env file)")
    parser.add_argument("--whisper-model", default="small", 
                       choices=["tiny", "base", "small", "medium", "large"],
                       help="Whisper model size (default: small)")
    parser.add_argument("--max-retries", type=int, default=3,
                       help="Maximum number of retry attempts (default: 3)")
    parser.add_argument("--env-file", default=".env",
                       help="Path to .env file (default: .env)")
    
    args = parser.parse_args()
    
    # Load custom .env file if specified
    if args.env_file != ".env":
        load_dotenv(args.env_file)
    
    # Get audio file path (interactive if not provided)
    audio_file_path = get_audio_file_path(args.audio_file)
    
    # Validate audio file
    if not Path(audio_file_path).exists():
        print(f"❌ Error: Audio file not found: {audio_file_path}")
        return 1
    
    # Load environment variables
    env_vars = load_environment_variables()
    if not env_vars:
        return 1
    
    # Use command line args if provided, otherwise use environment variables
    claude_api_key = args.claude_api_key or env_vars['ANTHROPIC_API_KEY']
    hf_token = args.hf_token or env_vars['HUGGINGFACE_TOKEN']
    
    try:
        print("🔑 Configuration loaded successfully")
        print(f"📁 Audio file: {audio_file_path}")
        print(f"🤖 Whisper model: {args.whisper_model}")
        print(f"🔄 Max retries: {args.max_retries}")
        print(f"🔑 Using API keys from: {'command line' if args.claude_api_key else '.env file'}")
        print()
        
        # Initialize pipeline
        pipeline = AudioTranscriptionPipeline(
            claude_api_key=claude_api_key,
            max_retries=args.max_retries
        )
        
        # Process audio file
        result = pipeline.process_audio_file(
            audio_file_path=audio_file_path,
            auth_token=hf_token,
            whisper_model_size=args.whisper_model
        )
        
        print(f"\n🎉 Processing complete!")
        print(f"📂 Results saved to: {result['output_dir']}")
        print(f"📊 Session ID: {result['session_id']}")
        
        return 0
        
    except Exception as e:
        print(f"❌ Error: {e}")
        import traceback
        print("\n📋 Full error traceback:")
        traceback.print_exc()
        return 1

if __name__ == "__main__":
    exit(main())